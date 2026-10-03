from io import BytesIO

import pymupdf
import pytest
from docx import Document

from personal_document_intelligence_api.documents.models import ExtractionMethod
from personal_document_intelligence_api.documents.ocr import OcrResult
from personal_document_intelligence_api.documents.parser import (
    UnsupportedDocumentTypeError,
    parse_document,
)


class FakeOcrEngine:
    def __init__(self) -> None:
        self.call_count = 0

    def recognize(self, image_bytes: bytes) -> OcrResult:
        self.call_count += 1

        return OcrResult(
            text="Text extracted using fake OCR",
            confidence=0.95,
        )


def test_parse_plain_text() -> None:
    document = parse_document(
        filename="notes.txt",
        file_bytes=b"Hello from the document",
        ocr_engine=FakeOcrEngine(),
    )

    assert document.file_type == "txt"
    assert document.text == "Hello from the document"
    assert document.native_section_count == 1
    assert document.ocr_section_count == 0


def test_parse_docx() -> None:
    docx = Document()
    docx.add_paragraph("Document heading")
    docx.add_paragraph("Important document content")

    buffer = BytesIO()
    docx.save(buffer)

    document = parse_document(
        filename="report.docx",
        file_bytes=buffer.getvalue(),
        ocr_engine=FakeOcrEngine(),
    )

    assert "Document heading" in document.text
    assert "Important document content" in document.text
    assert document.native_section_count == 1


def test_parse_native_pdf_without_ocr() -> None:
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text(
        (72, 72),
        "This is a native PDF containing enough text for direct extraction.",
    )
    pdf_bytes = pdf.tobytes()
    pdf.close()

    ocr_engine = FakeOcrEngine()

    document = parse_document(
        filename="native.pdf",
        file_bytes=pdf_bytes,
        ocr_engine=ocr_engine,
    )

    assert document.page_count == 1
    assert document.sections[0].extraction_method == ExtractionMethod.NATIVE
    assert ocr_engine.call_count == 0


def test_parse_scanned_pdf_uses_ocr() -> None:
    pdf = pymupdf.open()
    pdf.new_page()
    pdf_bytes = pdf.tobytes()
    pdf.close()

    ocr_engine = FakeOcrEngine()

    document = parse_document(
        filename="scanned.pdf",
        file_bytes=pdf_bytes,
        ocr_engine=ocr_engine,
    )

    assert document.page_count == 1
    assert document.sections[0].extraction_method == ExtractionMethod.OCR
    assert document.sections[0].confidence == 0.95
    assert ocr_engine.call_count == 1


def test_reject_unsupported_document_type() -> None:
    with pytest.raises(UnsupportedDocumentTypeError):
        parse_document(
            filename="spreadsheet.xlsx",
            file_bytes=b"content",
            ocr_engine=FakeOcrEngine(),
        )
