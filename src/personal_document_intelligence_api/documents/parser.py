from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document as open_docx_document

from .models import DocumentSection, ExtractionMethod, ParsedDocument
from .ocr import OcrEngine, OcrError

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt"}
MIN_NATIVE_CHARACTERS = 40
OCR_RENDER_SCALE = 3


class DocumentParsingError(Exception):
    """Raised when a document cannot be parsed."""


class UnsupportedDocumentTypeError(DocumentParsingError):
    """Raised when a document type is unsupported."""


class EmptyDocumentError(DocumentParsingError):
    """Raised when no readable text is found."""


def parse_document(
    filename: str,
    file_bytes: bytes,
    ocr_engine: OcrEngine,
) -> ParsedDocument:
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocumentTypeError(f"Unsupported document type: {extension or 'unknown'}")

    if not file_bytes:
        raise EmptyDocumentError("The document is empty")

    if extension == ".pdf":
        return _parse_pdf(filename, file_bytes, ocr_engine)

    if extension == ".docx":
        sections = _parse_docx(file_bytes)
        page_count = None
    else:
        sections = _parse_plain_text(file_bytes)
        page_count = None

    if not sections:
        raise EmptyDocumentError("No readable text was found")

    return ParsedDocument(
        filename=filename,
        file_type=extension.removeprefix("."),
        size_bytes=len(file_bytes),
        page_count=page_count,
        sections=sections,
    )


def _parse_pdf(
    filename: str,
    file_bytes: bytes,
    ocr_engine: OcrEngine,
) -> ParsedDocument:
    try:
        with pymupdf.open(stream=file_bytes, filetype="pdf") as document:
            if document.needs_pass:
                raise DocumentParsingError("Password-protected PDFs are not supported")

            sections: list[DocumentSection] = []

            for page_index, page in enumerate(document, start=1):
                native_text = _normalize_text(page.get_text("text"))

                if len(native_text) >= MIN_NATIVE_CHARACTERS:
                    sections.append(
                        DocumentSection(
                            text=native_text,
                            page_number=page_index,
                            extraction_method=ExtractionMethod.NATIVE,
                        )
                    )
                    continue

                pixmap = page.get_pixmap(
                    matrix=pymupdf.Matrix(
                        OCR_RENDER_SCALE,
                        OCR_RENDER_SCALE,
                    ),
                    alpha=False,
                )

                ocr_result = ocr_engine.recognize(pixmap.tobytes("png"))
                ocr_text = _normalize_text(ocr_result.text)

                if ocr_text:
                    sections.append(
                        DocumentSection(
                            text=ocr_text,
                            page_number=page_index,
                            extraction_method=ExtractionMethod.OCR,
                            confidence=ocr_result.confidence,
                        )
                    )
                elif native_text:
                    sections.append(
                        DocumentSection(
                            text=native_text,
                            page_number=page_index,
                            extraction_method=ExtractionMethod.NATIVE,
                        )
                    )

            page_count = len(document)

    except DocumentParsingError:
        raise
    except OcrError as error:
        raise DocumentParsingError(str(error)) from error
    except Exception as error:
        raise DocumentParsingError("The PDF could not be parsed") from error

    if not sections:
        raise EmptyDocumentError("No readable text was found in the PDF")

    return ParsedDocument(
        filename=filename,
        file_type="pdf",
        size_bytes=len(file_bytes),
        page_count=page_count,
        sections=tuple(sections),
    )


def _parse_docx(
    file_bytes: bytes,
) -> tuple[DocumentSection, ...]:
    try:
        document = open_docx_document(BytesIO(file_bytes))
        content: list[str] = []

        for paragraph in document.paragraphs:
            text = _normalize_text(paragraph.text)

            if text:
                content.append(text)

        for table in document.tables:
            for row in table.rows:
                cells: list[str] = []

                for cell in row.cells:
                    text = _normalize_text(cell.text)

                    if text:
                        cells.append(text)

                if cells:
                    content.append(" | ".join(cells))

    except Exception as error:
        raise DocumentParsingError("The DOCX document could not be parsed") from error

    text = "\n\n".join(content)

    if not text:
        return ()

    return (
        DocumentSection(
            text=text,
            extraction_method=ExtractionMethod.NATIVE,
        ),
    )


def _parse_plain_text(
    file_bytes: bytes,
) -> tuple[DocumentSection, ...]:
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError as error:
        raise DocumentParsingError("Text documents must use UTF-8 encoding") from error

    text = _normalize_text(text)

    if not text:
        return ()

    return (
        DocumentSection(
            text=text,
            extraction_method=ExtractionMethod.NATIVE,
        ),
    )


def _normalize_text(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()
